"""Unit tests for the fetch_document tool (extraction logic)."""

from __future__ import annotations

from io import BytesIO

from research_assistant.tools.general.fetch_document import (
    _detect_type,
    _extract_docx,
    _extract_pdf,
    _extract_pptx,
)

# ── Type detection ────────────────────────────────────────────────────────


def test_detect_pdf_by_content_type() -> None:
    assert _detect_type("https://example.com/file", "application/pdf") == "pdf"


def test_detect_pdf_by_extension() -> None:
    assert _detect_type("https://example.com/report.pdf", "application/octet-stream") == "pdf"


def test_detect_docx_by_content_type() -> None:
    ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert _detect_type("https://example.com/file", ct) == "docx"


def test_detect_docx_by_extension() -> None:
    assert _detect_type("https://example.com/doc.docx", "") == "docx"


def test_detect_pptx_by_extension() -> None:
    assert _detect_type("https://example.com/slides.pptx", "") == "pptx"


def test_detect_unknown_returns_none() -> None:
    assert _detect_type("https://example.com/file.xyz", "text/html") is None


def test_detect_ignores_query_string() -> None:
    assert _detect_type("https://example.com/report.pdf?token=abc", "") == "pdf"


# ── PDF extraction ────────────────────────────────────────────────────────


def test_extract_pdf() -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello from PDF page 1")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Page two content")
    data = doc.tobytes()
    doc.close()

    result = _extract_pdf(data)
    assert result["type"] == "pdf"
    assert result["page_count"] == 2
    assert "Hello from PDF" in result["pages"][0]["text"]
    assert result["pages"][1]["page"] == 2


# ── DOCX extraction ──────────────────────────────────────────────────────


def test_extract_docx() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    buf = BytesIO()
    doc.save(buf)

    result = _extract_docx(buf.getvalue())
    assert result["type"] == "docx"
    assert result["paragraph_count"] == 2
    assert "First paragraph" in result["paragraphs"]


# ── PPTX extraction ──────────────────────────────────────────────────────


def test_extract_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Slide One Title"
    slide.placeholders[1].text = "Bullet point here"

    slide2 = prs.slides.add_slide(layout)
    slide2.shapes.title.text = "Slide Two"

    buf = BytesIO()
    prs.save(buf)

    result = _extract_pptx(buf.getvalue())
    assert result["type"] == "pptx"
    assert result["slide_count"] == 2
    assert any("Slide One Title" in " ".join(s["text"]) for s in result["slides"])
