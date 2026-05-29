"""Manuscript + reviewer-response report (top-6 #5).

Renders the assembled IMRaD manuscript as a journal-ready PDF / DOCX.
Optionally appends the latest reviewer-response document so a single
download captures both artefacts when the user is mid-revision.

Pipeline mirrors the other report modules:

    Message rows ─▶ assemble_report_data() ─▶ ManuscriptReportData
                                                    │
                                                    ├─▶ build_pdf()  ─▶ bytes
                                                    └─▶ build_docx() ─▶ bytes

`images_dir` is kept for signature parity with the other builders but
unused — the manuscript carries no embedded figures (the source
meta-analysis report has the forest plots, downloaded separately).
"""

from __future__ import annotations

import io
import json
import logging
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from pydantic import ValidationError
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from research_assistant.domain.manuscript import (
    ManuscriptDraft,
    ManuscriptIntake,
    ReviewerResponse,
)
from research_assistant.persistence.models import Message
from research_assistant.reports._shared_styles import (
    BORDER,
    DOCX_MUTED,
    DOCX_NAVY,
    LIGHT,
    make_page_decorations,
    make_pdf_styles,
)

logger = logging.getLogger(__name__)


@dataclass
class ManuscriptReportData:
    thread_id: str
    title: str
    generated_at: datetime
    research_question: str
    intake: ManuscriptIntake | None
    draft: ManuscriptDraft | None
    reviewer_response: ReviewerResponse | None


def _peek_kind(final_answer_json: str) -> str | None:
    try:
        obj = json.loads(final_answer_json)
    except (json.JSONDecodeError, TypeError):
        return None
    if isinstance(obj, dict):
        kind = obj.get("kind")
        return str(kind) if isinstance(kind, str) else None
    return None


def assemble_report_data(
    thread_id: str,
    messages: Iterable[Message],
) -> ManuscriptReportData | None:
    """Walk an ordered message list and pull out the latest manuscript state.

    Returns None if the thread has no manuscript_draft turn yet — the
    intake + reviewer_response alone aren't useful as a downloadable
    artefact.
    """
    research_question = ""
    intake: ManuscriptIntake | None = None
    draft: ManuscriptDraft | None = None
    reviewer_response: ReviewerResponse | None = None

    for msg in messages:
        if msg.role == "user" and not research_question and msg.input_text:
            research_question = msg.input_text.strip()
            continue
        if msg.role != "assistant" or not msg.final_answer:
            continue
        kind = _peek_kind(msg.final_answer)
        try:
            if kind == "manuscript_intake":
                intake = ManuscriptIntake.model_validate_json(msg.final_answer)
            elif kind == "manuscript_draft":
                draft = ManuscriptDraft.model_validate_json(msg.final_answer)
            elif kind == "reviewer_response":
                reviewer_response = ReviewerResponse.model_validate_json(msg.final_answer)
        except ValidationError:
            logger.exception("Skipping malformed %s turn in message %s", kind, msg.id)

    if draft is None:
        return None

    return ManuscriptReportData(
        thread_id=thread_id,
        title=draft.title or "Manuscript",
        generated_at=datetime.now(UTC),
        research_question=research_question or "(Research question not captured.)",
        intake=intake,
        draft=draft,
        reviewer_response=reviewer_response,
    )


# ── PDF builder ──────────────────────────────────────────────────────────


def _abstract_block(draft: ManuscriptDraft, styles: dict[str, ParagraphStyle]) -> list[Any]:
    """Render the structured or free-form abstract as a Platypus flow."""
    out: list[Any] = []
    if draft.structured_abstract is not None:
        sa = draft.structured_abstract
        rows: list[list[Any]] = [
            [
                Paragraph("<b>Background</b>", styles["Cell"]),
                Paragraph(sa.background, styles["Cell"]),
            ],
            [Paragraph("<b>Methods</b>", styles["Cell"]), Paragraph(sa.methods, styles["Cell"])],
            [Paragraph("<b>Results</b>", styles["Cell"]), Paragraph(sa.results, styles["Cell"])],
            [
                Paragraph("<b>Conclusions</b>", styles["Cell"]),
                Paragraph(sa.conclusions, styles["Cell"]),
            ],
        ]
        t = Table(rows, colWidths=[1.3 * inch, 5.5 * inch])
        t.setStyle(
            TableStyle(
                [
                    ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT]),
                    ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                    ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        out.append(t)
    elif draft.abstract_text:
        out.append(Paragraph(draft.abstract_text, styles["Body"]))
    if draft.abstract_word_count:
        out.append(
            Paragraph(
                f"<i>Abstract word count: {draft.abstract_word_count}</i>",
                styles["Body"],
            )
        )
    return out


def _reference_paragraph(c: Any, styles: dict[str, ParagraphStyle]) -> Paragraph:
    meta_bits = [c.authors, c.journal, str(c.year) if c.year else None]
    meta_text = ", ".join(b for b in meta_bits if b)
    tail = ""
    if c.pmid:
        tail += f" PMID {c.pmid}"
    if c.doi:
        tail += f" doi:{c.doi}"
    if c.url:
        tail += f" — {c.url}"
    return Paragraph(f"[{c.n}] {c.title}. <i>{meta_text}</i>{tail}", styles["Body"])


def build_pdf(data: ManuscriptReportData, images_dir: Path) -> bytes:
    """Render the manuscript as a journal-style PDF.

    `images_dir` is unused — the manuscript itself carries no embedded
    figures. Forest plots / KM curves live in the upstream
    meta_analysis report, which downloads separately.
    """
    del images_dir  # signature-parity only
    buf = io.BytesIO()
    styles = make_pdf_styles()
    decor = make_page_decorations(footer_label="MANUSCRIPT DRAFT")
    doc = BaseDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.7 * inch,
        title=data.title,
    )
    frame = Frame(
        doc.leftMargin,
        doc.bottomMargin,
        doc.width,
        doc.height,
        id="main",
    )
    doc.addPageTemplates([PageTemplate(id="page", frames=[frame], onPage=decor)])

    flow: list[Any] = []
    draft = data.draft
    assert draft is not None  # assemble_report_data guarantees this
    journal = data.intake.journal_target if data.intake else "generic"

    flow.append(Paragraph(f"MANUSCRIPT DRAFT · target: {journal.upper()}", styles["Eyebrow"]))
    flow.append(Paragraph(draft.title, styles["Title"]))
    if draft.short_title:
        flow.append(Paragraph(f"<i>Running head:</i> {draft.short_title}", styles["Body"]))
    ts = data.generated_at.strftime("%Y-%m-%d %H:%M UTC")
    flow.append(Paragraph(f"Generated {ts}  ·  thread {data.thread_id[:8]}", styles["Body"]))
    flow.append(Spacer(1, 12))

    flow.append(Paragraph("Abstract", styles["H2"]))
    flow.extend(_abstract_block(draft, styles))

    flow.append(Spacer(1, 10))
    flow.append(Paragraph("Introduction", styles["H2"]))
    flow.append(Paragraph(draft.introduction, styles["Body"]))

    flow.append(Spacer(1, 10))
    flow.append(Paragraph("Methods", styles["H2"]))
    flow.append(Paragraph(draft.methods, styles["Body"]))

    flow.append(Spacer(1, 10))
    flow.append(Paragraph("Results", styles["H2"]))
    flow.append(Paragraph(draft.results, styles["Body"]))

    flow.append(Spacer(1, 10))
    flow.append(Paragraph("Discussion", styles["H2"]))
    flow.append(Paragraph(draft.discussion, styles["Body"]))

    if draft.references:
        flow.append(Spacer(1, 10))
        flow.append(Paragraph("References", styles["H2"]))
        for c in draft.references:
            flow.append(_reference_paragraph(c, styles))

    # Append the latest reviewer-response document when present.
    if data.reviewer_response is not None:
        rr = data.reviewer_response
        flow.append(Spacer(1, 16))
        flow.append(Paragraph("Reviewer response", styles["Eyebrow"]))
        flow.append(Paragraph("Point-by-point responses", styles["H2"]))
        if rr.cover_letter_text:
            flow.append(Paragraph(rr.cover_letter_text, styles["Body"]))
        for item in rr.responses:
            flow.append(Spacer(1, 6))
            flow.append(
                Paragraph(
                    f"<b>{item.reviewer_id}:</b> <i>{item.comment_excerpt}</i>",
                    styles["Body"],
                )
            )
            flow.append(Paragraph(item.response_text, styles["Body"]))
            if item.suggested_manuscript_edits:
                flow.append(
                    Paragraph(
                        f"<b>Suggested edit:</b> {item.suggested_manuscript_edits}",
                        styles["Body"],
                    )
                )
            if not item.is_addressed:
                flow.append(
                    Paragraph(
                        "<b>(Pushing back on this point — see response above.)</b>",
                        styles["Body"],
                    )
                )

    doc.build(flow)
    return buf.getvalue()


# ── DOCX builder ─────────────────────────────────────────────────────────


def build_docx(data: ManuscriptReportData, images_dir: Path) -> bytes:
    del images_dir  # signature-parity only
    docx = Document()
    draft = data.draft
    assert draft is not None
    journal = data.intake.journal_target if data.intake else "generic"

    docx.add_paragraph().add_run(f"MANUSCRIPT DRAFT · target: {journal.upper()}").bold = True
    title_p = docx.add_paragraph()
    title_run = title_p.add_run(draft.title)
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = DOCX_NAVY
    if draft.short_title:
        docx.add_paragraph(f"Running head: {draft.short_title}")
    ts = data.generated_at.strftime("%Y-%m-%d %H:%M UTC")
    meta_p = docx.add_paragraph()
    meta_p.add_run(f"Generated {ts}  ·  thread {data.thread_id[:8]}").font.color.rgb = DOCX_MUTED

    docx.add_heading("Abstract", level=2)
    if draft.structured_abstract is not None:
        sa = draft.structured_abstract
        rows = [
            ("Background", sa.background),
            ("Methods", sa.methods),
            ("Results", sa.results),
            ("Conclusions", sa.conclusions),
        ]
        tbl = docx.add_table(rows=len(rows), cols=2)
        tbl.style = "Light Grid Accent 1"
        for i, (k, v) in enumerate(rows):
            tbl.cell(i, 0).text = k
            tbl.cell(i, 1).text = v
    elif draft.abstract_text:
        docx.add_paragraph(draft.abstract_text)

    for heading, body in (
        ("Introduction", draft.introduction),
        ("Methods", draft.methods),
        ("Results", draft.results),
        ("Discussion", draft.discussion),
    ):
        docx.add_heading(heading, level=2)
        docx.add_paragraph(body)

    if draft.references:
        docx.add_heading("References", level=2)
        for c in draft.references:
            meta_bits = [c.authors, c.journal, str(c.year) if c.year else None]
            meta_text = ", ".join(b for b in meta_bits if b)
            tail = ""
            if c.pmid:
                tail += f" PMID {c.pmid}"
            if c.doi:
                tail += f" doi:{c.doi}"
            if c.url:
                tail += f" — {c.url}"
            docx.add_paragraph(f"[{c.n}] {c.title}. {meta_text}{tail}", style="List Bullet")

    if data.reviewer_response is not None:
        rr = data.reviewer_response
        docx.add_heading("Reviewer response", level=2)
        if rr.cover_letter_text:
            docx.add_paragraph(rr.cover_letter_text)
        for item in rr.responses:
            heading = docx.add_heading(level=3)
            heading.add_run(f"{item.reviewer_id}: ").bold = True
            heading.add_run(item.comment_excerpt).italic = True
            docx.add_paragraph(item.response_text)
            if item.suggested_manuscript_edits:
                docx.add_paragraph(f"Suggested edit: {item.suggested_manuscript_edits}")
            if not item.is_addressed:
                docx.add_paragraph("(Pushing back on this point.)")

    buf = io.BytesIO()
    docx.save(buf)
    return buf.getvalue()


__all__ = [
    "ManuscriptReportData",
    "assemble_report_data",
    "build_docx",
    "build_pdf",
]
