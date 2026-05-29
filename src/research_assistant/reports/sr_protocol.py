"""SR/MA protocol report — assemble a thread's most-recent ProtocolDocument
into a downloadable PRISMA-P-shaped PDF or DOCX.

Pipeline mirrors `reports/meta_analysis.py`:

    Message rows ─▶ assemble_report_data() ─▶ SrProtocolReportData
                                                    │
                                                    ├─▶ build_pdf()  ─▶ bytes
                                                    └─▶ build_docx() ─▶ bytes

Assembly pulls the latest ProtocolDocument; the embedded `methods` field
carries the full PICO + eligibility + RoB tool + synthesis plan, so no
walking back to earlier ProtocolMethods turns is needed.
"""

from __future__ import annotations

import io
import json
import logging
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from pydantic import ValidationError
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from research_assistant.domain.sr_protocol import (
    Citation,
    ProtocolDocument,
    ProtocolMethods,
)
from research_assistant.persistence.models import Message
from research_assistant.reports._shared_styles import (
    BORDER,
    DOCX_AMBER,
    DOCX_MUTED,
    DOCX_NAVY,
    DOCX_TEAL,
    LIGHT,
    NAVY,
    docx_set_heading_color,
    make_page_decorations,
    make_pdf_styles,
)

logger = logging.getLogger(__name__)

_USER_INPUT_PREFIX = "[USER INPUT NEEDED"


# ── Data assembly ────────────────────────────────────────────────────────


@dataclass
class SrProtocolReportData:
    """Everything the renderers need, pre-shaped from one thread.

    The full ProtocolMethods is carried inline so per-section renderers can
    reach in for PICO, eligibility, etc. without a second lookup.
    """

    thread_id: str
    research_question: str
    generated_at: datetime
    methods: ProtocolMethods
    background: str
    references: list[Citation]
    prospero_field_map: list[Any]  # ProsperoFieldMap shape (.field, .content)
    notes: str | None
    is_final: bool


def assemble_report_data(
    thread_id: str,
    messages: Iterable[Message],
) -> SrProtocolReportData | None:
    """Pull the latest ProtocolDocument from the thread's messages.

    Returns None if no protocol_document turn has been emitted yet — callers
    should map that to HTTP 404. We accept a non-final ProtocolDocument as
    well; the UI only shows the download buttons when `is_final` is true, so
    a non-final document arriving here means the user hit the URL directly.
    The `is_final` field on the dataclass surfaces that state to readers.
    """
    research_question = ""
    document: ProtocolDocument | None = None

    for msg in messages:
        if msg.role == "user" and not research_question and msg.input_text:
            research_question = msg.input_text.strip()
            continue
        if msg.role != "assistant" or not msg.final_answer:
            continue
        if _peek_kind(msg.final_answer) != "protocol_document":
            continue
        try:
            document = ProtocolDocument.model_validate_json(msg.final_answer)
        except ValidationError:
            logger.exception("Skipping malformed ProtocolDocument in message %s", msg.id)

    if document is None:
        return None

    return SrProtocolReportData(
        thread_id=thread_id,
        research_question=research_question or "(Research question not captured.)",
        generated_at=datetime.now(UTC),
        methods=document.methods,
        background=document.background,
        references=list(document.references),
        prospero_field_map=list(document.prospero_field_map),
        notes=document.notes,
        is_final=document.is_final,
    )


def _peek_kind(final_answer_json: str) -> str | None:
    try:
        obj = json.loads(final_answer_json)
    except (json.JSONDecodeError, TypeError):
        return None
    if isinstance(obj, dict):
        kind = obj.get("kind")
        return str(kind) if isinstance(kind, str) else None
    return None


# ── Shared rendering helpers ─────────────────────────────────────────────


def _format_origin(c: Citation) -> str:
    """Short provenance label per reference (mirrors the UI badge colours)."""
    suffix_parts: list[str] = []
    if c.pmid:
        suffix_parts.append(f"PMID {c.pmid}")
    if c.doi:
        suffix_parts.append(f"doi:{c.doi}")
    if c.url and not c.pmid and not c.doi:
        suffix_parts.append(c.url)
    suffix = "  ·  ".join(suffix_parts)
    return f"[{c.origin}]" + (f"  {suffix}" if suffix else "")


# ── PDF builder ──────────────────────────────────────────────────────────


def _build_kv_pdf_table(
    rows: list[tuple[str, str]], styles: dict[str, ParagraphStyle], label_w: float = 1.6
) -> Table:
    data: list[list[Any]] = [
        [
            Paragraph(f"<b>{label}</b>", styles["Cell"]),
            Paragraph(value or "—", styles["Cell"]),
        ]
        for label, value in rows
    ]
    t = Table(data, colWidths=[label_w * inch, (6.8 - label_w) * inch])
    t.setStyle(
        TableStyle(
            [
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT]),
                ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def _prospero_pdf_table(
    rows: list[Any], styles: dict[str, ParagraphStyle]
) -> Table | None:
    if not rows:
        return None
    data: list[list[Any]] = [
        [
            Paragraph("Field", styles["CellHead"]),
            Paragraph("Content", styles["CellHead"]),
        ]
    ]
    for row in rows:
        content = (row.content or "").strip()
        # User-input placeholders are flagged in amber so reviewers can spot
        # what they still owe before submitting to PROSPERO.
        cell_style = (
            styles["Amber"] if content.startswith(_USER_INPUT_PREFIX) else styles["Cell"]
        )
        data.append(
            [
                Paragraph(row.field or "—", styles["Cell"]),
                Paragraph(content or "—", cell_style),
            ]
        )
    t = Table(data, colWidths=[2.2 * inch, 4.6 * inch], repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
                ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def build_pdf(data: SrProtocolReportData, images_dir: Path) -> bytes:
    """Render a PRISMA-P-shaped PDF protocol. `images_dir` accepted for
    interface parity with the other report builders; not used here (the
    sr_protocol output has no embedded images today)."""
    del images_dir  # interface parity
    buf = io.BytesIO()
    styles = make_pdf_styles()
    doc = BaseDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=data.methods.title or "SR/MA protocol",
        author="Clinical Research Assistant",
    )
    w, h = LETTER
    doc.addPageTemplates(
        [
            PageTemplate(
                id="content",
                frames=[
                    Frame(
                        0.6 * inch,
                        0.7 * inch,
                        w - 1.2 * inch,
                        h - 1.5 * inch,
                        id="content",
                        showBoundary=0,
                    )
                ],
                onPage=make_page_decorations("SR/MA protocol"),
            )
        ]
    )

    story: list[Any] = []
    finality = "FINAL" if data.is_final else "DRAFT"
    story.append(
        Paragraph(
            f"PRISMA-P PROTOCOL  ·  {finality}  ·  "
            f"GENERATED {data.generated_at.strftime('%Y-%m-%d %H:%M UTC')}",
            styles["Eyebrow"],
        )
    )
    story.append(Paragraph(data.methods.title or "SR/MA protocol", styles["Title"]))

    story.append(Paragraph("Research question", styles["H2"]))
    story.append(Paragraph(data.research_question, styles["Quote"]))

    # ── Methodological core ──────────────────────────────────────────────
    story.append(Paragraph("Objectives (PICO)", styles["H2"]))
    pico = data.methods.pico
    pico_rows: list[tuple[str, str]] = [
        ("Population", pico.population),
        ("Intervention", pico.intervention),
        ("Comparison", pico.comparison),
        ("Outcomes", "; ".join(pico.outcomes) or "—"),
        ("Study designs", "; ".join(pico.study_types) or "—"),
    ]
    if pico.age_range:
        pico_rows.append(("Age range", pico.age_range))
    story.append(_build_kv_pdf_table(pico_rows, styles))

    story.append(Paragraph(f"Review type: {data.methods.review_type}", styles["H3"]))

    story.append(Paragraph("Eligibility", styles["H2"]))
    eligibility = data.methods.eligibility
    elig_rows: list[tuple[str, str]] = [
        ("Inclusion", "; ".join(eligibility.inclusion) or "—"),
        ("Exclusion", "; ".join(eligibility.exclusion) or "—"),
        ("Study designs", "; ".join(eligibility.study_designs) or "—"),
        ("Language", "; ".join(eligibility.language) or "—"),
        ("Date range", eligibility.date_range or "—"),
        ("Age range", eligibility.age_range or "—"),
        ("Setting", eligibility.setting or "—"),
    ]
    story.append(_build_kv_pdf_table(elig_rows, styles))

    story.append(Paragraph("Information sources", styles["H2"]))
    story.append(
        Paragraph(
            "; ".join(data.methods.information_sources) or "—",
            styles["Body"],
        )
    )

    story.append(Paragraph("Risk-of-bias plan", styles["H2"]))
    story.append(
        Paragraph(
            f"<b>Tool:</b> {data.methods.rob_tool.tool}", styles["Body"]
        )
    )
    story.append(
        Paragraph(data.methods.rob_tool.rationale or "—", styles["Body"])
    )

    if data.methods.effect_measures_plan:
        story.append(Paragraph("Effect measures plan", styles["H2"]))
        em_rows = [
            (outcome, measure)
            for outcome, measure in data.methods.effect_measures_plan.items()
        ]
        story.append(
            _build_kv_pdf_table(
                [(o, str(m)) for o, m in em_rows], styles, label_w=3.0
            )
        )

    story.append(Paragraph("Synthesis plan", styles["H2"]))
    sp = data.methods.synthesis_plan
    sp_rows: list[tuple[str, str]] = [
        ("Primary method", sp.primary_method),
        ("Heterogeneity", "; ".join(sp.heterogeneity_assessment) or "—"),
        ("Planned subgroups", "; ".join(sp.planned_subgroups) or "—"),
        (
            "Sensitivity analyses",
            "; ".join(sp.planned_sensitivity_analyses) or "—",
        ),
        ("Publication bias", "; ".join(sp.publication_bias_methods) or "—"),
        ("GRADE", "yes" if data.methods.use_grade else "no"),
    ]
    story.append(_build_kv_pdf_table(sp_rows, styles))

    # ── Background ───────────────────────────────────────────────────────
    story.append(Paragraph("Background", styles["H2"]))
    for para in [p.strip() for p in data.background.split("\n\n") if p.strip()]:
        story.append(Paragraph(para, styles["Body"]))

    # ── References ──────────────────────────────────────────────────────
    if data.references:
        story.append(Paragraph("References", styles["H2"]))
        for i, ref in enumerate(data.references, start=1):
            story.append(
                KeepTogether(
                    [
                        Paragraph(f"<b>[{i}]</b> {ref.text}", styles["Body"]),
                        Paragraph(_format_origin(ref), styles["Het"]),
                    ]
                )
            )

    # ── PROSPERO field map ──────────────────────────────────────────────
    pf_tbl = _prospero_pdf_table(data.prospero_field_map, styles)
    if pf_tbl is not None:
        story.append(Paragraph("PROSPERO registration field map", styles["H2"]))
        story.append(
            Paragraph(
                "Amber rows are placeholders for user-supplied content; "
                "fill them in before submitting to PROSPERO.",
                styles["Het"],
            )
        )
        story.append(pf_tbl)

    if data.notes:
        story.append(Paragraph("Notes", styles["H2"]))
        story.append(Paragraph(data.notes, styles["Body"]))

    story.append(Spacer(1, 0.1 * inch))
    if not data.is_final:
        story.append(
            Paragraph(
                "<i>Draft — not finalised by the reviewer.</i>", styles["Het"]
            )
        )

    doc.build(story)
    return buf.getvalue()


# ── DOCX builder ─────────────────────────────────────────────────────────


def _docx_kv_table(doc: Any, rows: list[tuple[str, str]]) -> None:
    if not rows:
        return
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for label, value in rows:
        row = t.add_row().cells
        row[0].text = label
        row[1].text = value or "—"
        if row[0].paragraphs[0].runs:
            row[0].paragraphs[0].runs[0].bold = True


def build_docx(data: SrProtocolReportData, images_dir: Path) -> bytes:
    """Render a PRISMA-P-shaped DOCX protocol. `images_dir` unused (interface
    parity with other report builders)."""
    del images_dir
    doc = Document()

    eyebrow = doc.add_paragraph()
    finality = "FINAL" if data.is_final else "DRAFT"
    eyebrow_run = eyebrow.add_run(
        f"PRISMA-P protocol  ·  {finality}  ·  "
        f"Generated {data.generated_at.strftime('%Y-%m-%d %H:%M UTC')}"
    )
    eyebrow_run.italic = True
    eyebrow_run.font.size = Pt(9)
    eyebrow_run.font.color.rgb = DOCX_MUTED

    title = doc.add_heading(data.methods.title or "SR/MA protocol", level=0)
    docx_set_heading_color(title, DOCX_NAVY)

    # Research question
    h = doc.add_heading("Research question", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    p = doc.add_paragraph(data.research_question)
    for run in p.runs:
        run.italic = True

    # PICO
    h = doc.add_heading("Objectives (PICO)", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    pico = data.methods.pico
    pico_rows: list[tuple[str, str]] = [
        ("Population", pico.population),
        ("Intervention", pico.intervention),
        ("Comparison", pico.comparison),
        ("Outcomes", "; ".join(pico.outcomes) or "—"),
        ("Study designs", "; ".join(pico.study_types) or "—"),
    ]
    if pico.age_range:
        pico_rows.append(("Age range", pico.age_range))
    _docx_kv_table(doc, pico_rows)

    sub = doc.add_heading(f"Review type: {data.methods.review_type}", level=2)
    docx_set_heading_color(sub, DOCX_NAVY)

    # Eligibility
    h = doc.add_heading("Eligibility", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    eligibility = data.methods.eligibility
    _docx_kv_table(
        doc,
        [
            ("Inclusion", "; ".join(eligibility.inclusion) or "—"),
            ("Exclusion", "; ".join(eligibility.exclusion) or "—"),
            ("Study designs", "; ".join(eligibility.study_designs) or "—"),
            ("Language", "; ".join(eligibility.language) or "—"),
            ("Date range", eligibility.date_range or "—"),
            ("Age range", eligibility.age_range or "—"),
            ("Setting", eligibility.setting or "—"),
        ],
    )

    # Information sources
    h = doc.add_heading("Information sources", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    doc.add_paragraph("; ".join(data.methods.information_sources) or "—")

    # Risk-of-bias
    h = doc.add_heading("Risk-of-bias plan", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    p = doc.add_paragraph()
    p.add_run("Tool: ").bold = True
    p.add_run(data.methods.rob_tool.tool)
    doc.add_paragraph(data.methods.rob_tool.rationale or "—")

    # Effect measures
    if data.methods.effect_measures_plan:
        h = doc.add_heading("Effect measures plan", level=1)
        docx_set_heading_color(h, DOCX_TEAL)
        _docx_kv_table(
            doc,
            [
                (outcome, str(measure))
                for outcome, measure in data.methods.effect_measures_plan.items()
            ],
        )

    # Synthesis plan
    h = doc.add_heading("Synthesis plan", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    sp = data.methods.synthesis_plan
    _docx_kv_table(
        doc,
        [
            ("Primary method", sp.primary_method),
            ("Heterogeneity", "; ".join(sp.heterogeneity_assessment) or "—"),
            ("Planned subgroups", "; ".join(sp.planned_subgroups) or "—"),
            (
                "Sensitivity analyses",
                "; ".join(sp.planned_sensitivity_analyses) or "—",
            ),
            ("Publication bias", "; ".join(sp.publication_bias_methods) or "—"),
            ("GRADE", "yes" if data.methods.use_grade else "no"),
        ],
    )

    # Background
    h = doc.add_heading("Background", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    for para in [p.strip() for p in data.background.split("\n\n") if p.strip()]:
        doc.add_paragraph(para)

    # References
    if data.references:
        h = doc.add_heading("References", level=1)
        docx_set_heading_color(h, DOCX_TEAL)
        # "List Number" style auto-numbers — no manual index needed.
        for ref in data.references:
            p = doc.add_paragraph(style="List Number")
            p.add_run(ref.text)
            p.add_run(f"  {_format_origin(ref)}").italic = True

    # PROSPERO field map
    if data.prospero_field_map:
        h = doc.add_heading("PROSPERO registration field map", level=1)
        docx_set_heading_color(h, DOCX_TEAL)
        note = doc.add_paragraph(
            "Amber rows are placeholders for user-supplied content; fill them "
            "in before submitting to PROSPERO."
        )
        for run in note.runs:
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = DOCX_MUTED
        t = doc.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Content"
        for cell in hdr:
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in data.prospero_field_map:
            cells = t.add_row().cells
            cells[0].text = row.field or "—"
            content = row.content or "—"
            cells[1].text = content
            if content.startswith(_USER_INPUT_PREFIX):
                for run in cells[1].paragraphs[0].runs:
                    run.italic = True
                    run.font.color.rgb = DOCX_AMBER

    if data.notes:
        h = doc.add_heading("Notes", level=1)
        docx_set_heading_color(h, DOCX_TEAL)
        doc.add_paragraph(data.notes)

    if not data.is_final:
        p = doc.add_paragraph("Draft — not finalised by the reviewer.")
        for run in p.runs:
            run.italic = True
            run.font.color.rgb = DOCX_MUTED

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
