"""Meta-analysis report — assemble persisted thread state into a manuscript-
style PDF or DOCX a researcher can take into a journal submission, IRB form,
or grant application.

Pipeline:

    Message rows ─▶ assemble_report_data() ─▶ MetaAnalysisReportData
                                                    │
                                                    ├─▶ build_pdf()  ─▶ bytes
                                                    └─▶ build_docx() ─▶ bytes

The assembler walks an ordered list of `Message` rows for one thread. The
two builders consume the pure dataclass — no DB access — so they're trivially
unit-testable. The endpoint glue lives in `web/threads.py`.

Image policy: a forest plot's `forest_plot_image` field is a URL path like
`/images/<uuid>_*.png`. The builders resolve that to a local file under
`images_dir` (= `settings.images_dir`, mounted into the agent container).
If the file is missing on disk, the report renders "Forest plot file
unavailable" in place of the image — never crashes.
"""

from __future__ import annotations

import io
import logging
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Inches, Pt
from pydantic import ValidationError
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus import (
    Image as RLImage,
)

from research_assistant.domain.meta_analysis import (
    DataExtraction,
    MetaAnalysisOutcomeResult,
    MetaAnalysisResults,
    PicoDraft,
    PicoTable,
    StudyExtractedData,
)
from research_assistant.persistence.models import Message
from research_assistant.reports._shared_styles import (
    BORDER,
    DOCX_MUTED,
    DOCX_NAVY,
    DOCX_TEAL,
    LIGHT,
    NAVY,
    docx_set_heading_color,
    make_page_decorations,
    make_pdf_styles,
)

logger = logging.getLogger(__name__)


# ── Data assembly ────────────────────────────────────────────────────────


@dataclass
class MetaAnalysisReportData:
    """Everything the renderers need, pre-shaped from one thread."""

    thread_id: str
    title: str
    research_question: str
    generated_at: datetime
    pico: PicoTable | None
    included_studies: list[StudyExtractedData]
    excluded_studies: list[dict[str, str]]
    studies_included_pmids: list[str]
    outcome_results: list[MetaAnalysisOutcomeResult]
    summary: str
    caveats: list[str] = field(default_factory=list)


def assemble_report_data(
    thread_id: str,
    messages: Iterable[Message],
) -> MetaAnalysisReportData | None:
    """Walk an ordered message list and pull out the latest meta-analysis state.

    Returns None if the thread never reached a meta-analysis turn — callers
    should map that to HTTP 404.

    Earlier card kinds (`pico`, `data_extraction`) are also pulled out so the
    report can show the PICO and the included studies table. Only the LAST
    instance of each kind is used; if the user re-confirmed PICO or re-ran
    extraction, the most recent edit wins.
    """
    research_question = ""
    pico: PicoTable | None = None
    extraction: DataExtraction | None = None
    analysis: MetaAnalysisResults | None = None

    for msg in messages:
        if msg.role == "user" and not research_question and msg.input_text:
            research_question = msg.input_text.strip()
            continue
        if msg.role != "assistant" or not msg.final_answer:
            continue
        # final_answer is a model_dump_json() of one of the union variants.
        # Parse each variant tolerantly — discriminator='kind' is on every
        # shape but pydantic_ai's validator wrapping changes between versions.
        kind = _peek_kind(msg.final_answer)
        if kind == "pico":
            try:
                pico = PicoDraft.model_validate_json(msg.final_answer).pico
            except ValidationError:
                logger.exception("Skipping malformed PicoDraft in message %s", msg.id)
        elif kind == "data_extraction":
            try:
                extraction = DataExtraction.model_validate_json(msg.final_answer)
            except ValidationError:
                logger.exception(
                    "Skipping malformed DataExtraction in message %s", msg.id
                )
        elif kind == "meta_analysis":
            try:
                analysis = MetaAnalysisResults.model_validate_json(msg.final_answer)
            except ValidationError:
                logger.exception(
                    "Skipping malformed MetaAnalysisResults in message %s", msg.id
                )

    if analysis is None:
        return None

    title = "Meta-analysis report"
    if pico and pico.intervention and pico.population:
        title = f"Meta-analysis — {pico.intervention} for {pico.population}"

    return MetaAnalysisReportData(
        thread_id=thread_id,
        title=title,
        research_question=research_question or "(Research question not captured.)",
        generated_at=datetime.now(UTC),
        pico=pico,
        included_studies=list(extraction.studies) if extraction else [],
        excluded_studies=list(analysis.studies_excluded),
        studies_included_pmids=list(analysis.studies_included),
        outcome_results=list(analysis.outcome_results),
        summary=analysis.summary,
        caveats=list(analysis.caveats),
    )


def _peek_kind(final_answer_json: str) -> str | None:
    """Cheap discriminator read without full validation."""
    import json

    try:
        obj = json.loads(final_answer_json)
    except (json.JSONDecodeError, TypeError):
        return None
    if isinstance(obj, dict):
        kind = obj.get("kind")
        return str(kind) if isinstance(kind, str) else None
    return None


# ── Image resolution ─────────────────────────────────────────────────────


def _resolve_image_path(url_path: str | None, images_dir: Path) -> Path | None:
    """`/images/abc_plot.png` → `<images_dir>/abc_plot.png` if the file exists."""
    if not url_path or not url_path.startswith("/images/"):
        return None
    candidate = images_dir / url_path.removeprefix("/images/")
    return candidate if candidate.is_file() else None


# ── Effect-size formatter (shared) ───────────────────────────────────────


def _format_effect(r: MetaAnalysisOutcomeResult) -> str:
    return (
        f"{r.effect_measure} = {r.pooled_effect:.2f} "
        f"(95% CI {r.ci_lower:.2f}–{r.ci_upper:.2f})"
    )


def _format_heterogeneity(r: MetaAnalysisOutcomeResult) -> str:
    parts: list[str] = []
    if r.i_squared is not None:
        parts.append(f"I² = {r.i_squared:.1f}%")
    if r.heterogeneity_p is not None:
        parts.append(f"heterogeneity p = {r.heterogeneity_p:.3f}")
    if r.p_value is not None:
        parts.append(f"p = {r.p_value:.3f}")
    parts.append(f"studies = {r.n_studies}")
    parts.append(f"n = {r.n_participants:,}")
    return "  ·  ".join(parts)


# ── PDF builder ──────────────────────────────────────────────────────────


def _pico_pdf_table(pico: PicoTable, styles: dict[str, ParagraphStyle]) -> Table:
    rows: list[list[Any]] = [
        [
            Paragraph("<b>Population</b>", styles["Cell"]),
            Paragraph(pico.population or "—", styles["Cell"]),
        ],
        [
            Paragraph("<b>Intervention</b>", styles["Cell"]),
            Paragraph(pico.intervention or "—", styles["Cell"]),
        ],
        [
            Paragraph("<b>Comparison</b>", styles["Cell"]),
            Paragraph(pico.comparison or "—", styles["Cell"]),
        ],
        [
            Paragraph("<b>Outcomes</b>", styles["Cell"]),
            Paragraph("; ".join(pico.outcomes) or "—", styles["Cell"]),
        ],
        [
            Paragraph("<b>Study designs</b>", styles["Cell"]),
            Paragraph("; ".join(pico.study_types) or "—", styles["Cell"]),
        ],
    ]
    if pico.age_range:
        rows.append(
            [
                Paragraph("<b>Age range</b>", styles["Cell"]),
                Paragraph(pico.age_range, styles["Cell"]),
            ]
        )
    t = Table(rows, colWidths=[1.4 * inch, 5.4 * inch])
    t.setStyle(
        TableStyle(
            [
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT]),
                ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def _studies_pdf_table(
    studies: list[StudyExtractedData], styles: dict[str, ParagraphStyle]
) -> Table | None:
    if not studies:
        return None
    headers = ["PMID", "Title", "Design", "Follow-up"]
    rows: list[list[Any]] = [
        [Paragraph(h, styles["CellHead"]) for h in headers]
    ]
    for s in studies:
        rows.append(
            [
                Paragraph(s.pmid or "—", styles["Cell"]),
                Paragraph(s.title or "—", styles["Cell"]),
                Paragraph(s.study_design or "—", styles["Cell"]),
                Paragraph(s.follow_up or "—", styles["Cell"]),
            ]
        )
    t = Table(
        rows,
        colWidths=[0.8 * inch, 4.0 * inch, 1.4 * inch, 0.9 * inch],
        repeatRows=1,
    )
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
                ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def build_pdf(data: MetaAnalysisReportData, images_dir: Path) -> bytes:
    """Render a manuscript-style PDF report. Returns the file bytes."""
    buf = io.BytesIO()
    styles = make_pdf_styles()
    doc = BaseDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=data.title,
        author="Clinical Research Assistant",
    )
    w, h = LETTER
    doc.addPageTemplates(
        [
            PageTemplate(
                id="content",
                frames=[
                    Frame(
                        0.6 * inch,
                        0.7 * inch,
                        w - 1.2 * inch,
                        h - 1.5 * inch,
                        id="content",
                        showBoundary=0,
                    )
                ],
                onPage=make_page_decorations("Meta-analysis report"),
            )
        ]
    )

    story: list[Any] = []
    story.append(
        Paragraph(
            f"GENERATED {data.generated_at.strftime('%Y-%m-%d %H:%M UTC')}",
            styles["Eyebrow"],
        )
    )
    story.append(Paragraph(data.title, styles["Title"]))

    story.append(Paragraph("Research question", styles["H2"]))
    story.append(Paragraph(data.research_question, styles["Quote"]))

    if data.pico:
        story.append(Paragraph("PICO", styles["H2"]))
        story.append(_pico_pdf_table(data.pico, styles))

    story.append(Paragraph("Included studies", styles["H2"]))
    studies_tbl = _studies_pdf_table(data.included_studies, styles)
    if studies_tbl is not None:
        story.append(studies_tbl)
    else:
        story.append(
            Paragraph(
                f"{len(data.studies_included_pmids)} studies were pooled. "
                "Per-study extraction details are not available in this thread.",
                styles["Body"],
            )
        )

    if data.excluded_studies:
        story.append(Paragraph("Excluded from analysis", styles["H3"]))
        excluded_text = "; ".join(
            f"{e.get('pmid', '?')} ({e.get('reason', 'no reason given')})"
            for e in data.excluded_studies
        )
        story.append(Paragraph(excluded_text, styles["Body"]))

    story.append(Paragraph("Results", styles["H2"]))
    for r in data.outcome_results:
        outcome_block: list[Any] = [
            Paragraph(f"{r.outcome} ({r.effect_measure})", styles["H3"]),
            Paragraph(_format_effect(r), styles["Stat"]),
            Paragraph(_format_heterogeneity(r), styles["Het"]),
        ]
        local_image = _resolve_image_path(r.forest_plot_image, images_dir)
        if local_image is not None:
            # Cap width at 6.4" and compute height from the image's natural
            # aspect ratio so the embed never distorts and never exceeds the
            # 7.3"-wide content frame.
            try:
                from PIL import Image as PILImage

                with PILImage.open(local_image) as im:
                    iw, ih = im.size
                max_w = 6.4 * inch
                draw_w = max_w
                draw_h = max_w * (ih / iw) if iw else 4.0 * inch
                outcome_block.append(
                    RLImage(str(local_image), width=draw_w, height=draw_h)
                )
            except Exception as exc:  # pragma: no cover - defensive
                logger.warning("Failed to embed forest plot %s: %s", local_image, exc)
                outcome_block.append(
                    Paragraph(
                        "<i>Forest plot could not be embedded.</i>", styles["Body"]
                    )
                )
        elif r.forest_plot_image:
            outcome_block.append(
                Paragraph(
                    "<i>Forest plot file unavailable on this server.</i>",
                    styles["Body"],
                )
            )
        outcome_block.append(Paragraph(r.interpretation, styles["Quote"]))
        story.append(KeepTogether(outcome_block))
        story.append(Spacer(1, 0.08 * inch))

    story.append(Paragraph("Summary", styles["H2"]))
    story.append(Paragraph(data.summary, styles["Body"]))

    if data.caveats:
        story.append(Paragraph("Caveats", styles["H2"]))
        for c in data.caveats:
            story.append(Paragraph(f"• {c}", styles["Body"]))

    doc.build(story)
    return buf.getvalue()


# ── DOCX builder ─────────────────────────────────────────────────────────


def build_docx(data: MetaAnalysisReportData, images_dir: Path) -> bytes:
    """Render a manuscript-style DOCX report. Returns the file bytes.

    Uses python-docx's built-in styles so the output opens cleanly in Word,
    LibreOffice, and Google Docs without further styling.
    """
    doc = Document()

    # Eyebrow line + title
    eyebrow = doc.add_paragraph()
    eyebrow_run = eyebrow.add_run(
        f"Generated {data.generated_at.strftime('%Y-%m-%d %H:%M UTC')}"
    )
    eyebrow_run.italic = True
    eyebrow_run.font.size = Pt(9)
    eyebrow_run.font.color.rgb = DOCX_MUTED

    title = doc.add_heading(data.title, level=0)
    docx_set_heading_color(title, DOCX_NAVY)

    # Research question
    h = doc.add_heading("Research question", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    p = doc.add_paragraph(data.research_question)
    for run in p.runs:
        run.italic = True

    # PICO table
    if data.pico:
        h = doc.add_heading("PICO", level=1)
        docx_set_heading_color(h, DOCX_TEAL)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for label, value in [
            ("Population", data.pico.population),
            ("Intervention", data.pico.intervention),
            ("Comparison", data.pico.comparison),
            ("Outcomes", "; ".join(data.pico.outcomes) or "—"),
            ("Study designs", "; ".join(data.pico.study_types) or "—"),
        ]:
            row = table.add_row().cells
            row[0].text = label
            row[1].text = value or "—"
            row[0].paragraphs[0].runs[0].bold = True
        if data.pico.age_range:
            row = table.add_row().cells
            row[0].text = "Age range"
            row[1].text = data.pico.age_range
            row[0].paragraphs[0].runs[0].bold = True

    # Included studies
    h = doc.add_heading("Included studies", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    if data.included_studies:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "PMID"
        hdr[1].text = "Title"
        hdr[2].text = "Design"
        hdr[3].text = "Follow-up"
        for cell in hdr:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for s in data.included_studies:
            row = table.add_row().cells
            row[0].text = s.pmid or "—"
            row[1].text = s.title or "—"
            row[2].text = s.study_design or "—"
            row[3].text = s.follow_up or "—"
    else:
        doc.add_paragraph(
            f"{len(data.studies_included_pmids)} studies were pooled. "
            "Per-study extraction details are not available in this thread."
        )

    if data.excluded_studies:
        h = doc.add_heading("Excluded from analysis", level=2)
        docx_set_heading_color(h, DOCX_NAVY)
        doc.add_paragraph(
            "; ".join(
                f"{e.get('pmid', '?')} ({e.get('reason', 'no reason given')})"
                for e in data.excluded_studies
            )
        )

    # Results
    h = doc.add_heading("Results", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    for r in data.outcome_results:
        sub = doc.add_heading(f"{r.outcome} ({r.effect_measure})", level=2)
        docx_set_heading_color(sub, DOCX_NAVY)
        eff = doc.add_paragraph()
        eff_run = eff.add_run(_format_effect(r))
        eff_run.bold = True
        eff_run.font.color.rgb = DOCX_NAVY
        het = doc.add_paragraph(_format_heterogeneity(r))
        for run in het.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = DOCX_MUTED
        local_image = _resolve_image_path(r.forest_plot_image, images_dir)
        if local_image is not None:
            try:
                doc.add_picture(str(local_image), width=Inches(6.4))
            except Exception as exc:  # pragma: no cover - defensive
                logger.warning("Failed to embed forest plot %s: %s", local_image, exc)
                err = doc.add_paragraph("Forest plot could not be embedded.")
                for run in err.runs:
                    run.italic = True
        elif r.forest_plot_image:
            note = doc.add_paragraph(
                "Forest plot file unavailable on this server."
            )
            for run in note.runs:
                run.italic = True
        interp = doc.add_paragraph(r.interpretation)
        for run in interp.runs:
            run.italic = True

    # Summary
    h = doc.add_heading("Summary", level=1)
    docx_set_heading_color(h, DOCX_TEAL)
    doc.add_paragraph(data.summary)

    # Caveats
    if data.caveats:
        h = doc.add_heading("Caveats", level=1)
        docx_set_heading_color(h, DOCX_TEAL)
        for c in data.caveats:
            doc.add_paragraph(c, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
