"""NMA report — league table + SUCRA + network geometry + PDF/DOCX.

Pipeline mirrors the other report modules:

    Message rows ─▶ assemble_report_data() ─▶ NmaReportData
                                                  │
                                                  ├─▶ build_pdf()  ─▶ bytes
                                                  └─▶ build_docx() ─▶ bytes

PDF is landscape (the league table is square; SUCRA needs horizontal
room). Embeds the sandbox-rendered network-geometry PNG inline when
`NetworkGraph.image_url` resolves to a file in `images_dir`.
"""

from __future__ import annotations

import io
import json
import logging
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches
from pydantic import ValidationError
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from research_assistant.domain.nma import (
    NmaResults,
)
from research_assistant.persistence.models import Message
from research_assistant.reports._shared_styles import (
    BORDER,
    DOCX_MUTED,
    LIGHT,
    make_page_decorations,
    make_pdf_styles,
)

logger = logging.getLogger(__name__)


@dataclass
class NmaReportData:
    thread_id: str
    title: str
    generated_at: datetime
    results: NmaResults | None


def _peek_kind(s: str) -> str | None:
    try:
        kind = json.loads(s).get("kind")
    except json.JSONDecodeError:
        return None
    return str(kind) if kind is not None else None


def assemble_report_data(
    thread_id: str,
    messages: Iterable[Message],
) -> NmaReportData | None:
    results: NmaResults | None = None
    for msg in messages:
        if msg.role != "assistant" or not msg.final_answer:
            continue
        if _peek_kind(msg.final_answer) != "nma_results":
            continue
        try:
            results = NmaResults.model_validate_json(msg.final_answer)
        except ValidationError:
            logger.exception("Skipping malformed nma_results in message %s", msg.id)
    if results is None:
        return None
    title = (
        f"NMA — {results.pico.outcome[:60]}"
        if results.pico.outcome
        else "Network meta-analysis"
    )
    return NmaReportData(
        thread_id=thread_id,
        title=title,
        generated_at=datetime.now(UTC),
        results=results,
    )


# ── Helpers ─────────────────────────────────────────────────────────────


def _resolve_image(images_dir: Path | None, url: str | None) -> Path | None:
    if images_dir is None or not url:
        return None
    name = url
    if name.startswith("/images/"):
        name = name[len("/images/") :]
    candidate = images_dir / name
    return candidate if candidate.exists() else None


def _fmt_effect(value: float | None, lo: float | None, hi: float | None) -> str:
    if value is None or lo is None or hi is None:
        return "—"
    return f"{value:.2f} ({lo:.2f}–{hi:.2f})"


def _direction(row_effect: float, effect_measure: str, ci_lo: float, ci_hi: float) -> str:
    if effect_measure in ("OR", "RR", "HR"):
        if ci_lo > 1.0:
            return "above"
        if ci_hi < 1.0:
            return "below"
        return "null"
    if ci_lo > 0:
        return "above"
    if ci_hi < 0:
        return "below"
    return "null"


def _league_pdf(results: NmaResults, styles: dict[str, ParagraphStyle]) -> Table:
    """Render the league table as a square matrix; coloured by direction."""
    interventions = results.pico.interventions
    by_pair: dict[tuple[str, str], Any] = {}
    for league_row in results.league_table.rows:
        by_pair[(league_row.row_intervention, league_row.col_intervention)] = league_row
    header = [Paragraph(f"<b>{c}</b>", styles["Cell"]) for c in [""] + list(interventions)]
    rows: list[list[Paragraph]] = [header]
    chip_cells: list[tuple[int, int, colors.Color]] = []
    for i, row_name in enumerate(interventions, start=1):
        cells: list[Paragraph] = [Paragraph(f"<b>{row_name}</b>", styles["Cell"])]
        for j, col_name in enumerate(interventions, start=1):
            if row_name == col_name:
                cells.append(Paragraph("—", styles["Cell"]))
                chip_cells.append((i, j, colors.HexColor("#e9ecef")))
                continue
            cell = by_pair.get((row_name, col_name))
            if cell is None:
                cells.append(Paragraph("—", styles["Cell"]))
                continue
            direction = _direction(
                cell.effect,
                results.league_table.effect_measure,
                cell.ci_lower,
                cell.ci_upper,
            )
            fill = {
                "above": colors.HexColor("#f8d7da"),  # risk
                "below": colors.HexColor("#d4edda"),  # protective
                "null": colors.white,
            }.get(direction, colors.white)
            cells.append(
                Paragraph(
                    _fmt_effect(cell.effect, cell.ci_lower, cell.ci_upper),
                    styles["Cell"],
                )
            )
            chip_cells.append((i, j, fill))
        rows.append(cells)
    n = len(interventions)
    width = max(1.0, 7.5 / (n + 1))
    t = Table(rows, colWidths=[width * inch] * (n + 1), repeatRows=1)
    style_cmds: list[Any] = [
        ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
        ("BACKGROUND", (0, 1), (0, -1), LIGHT),
        ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for row, col, fill in chip_cells:
        style_cmds.append(("BACKGROUND", (col, row), (col, row), fill))
    t.setStyle(TableStyle(style_cmds))
    return t


def _sucra_pdf(results: NmaResults, styles: dict[str, ParagraphStyle]) -> Table:
    header = [
        Paragraph(f"<b>{c}</b>", styles["Cell"])
        for c in ("Intervention", "Rank", "SUCRA", "Mean rank")
    ]
    sorted_rows = sorted(results.sucra, key=lambda r: r.rank)
    body: list[list[Paragraph]] = [header]
    for r in sorted_rows:
        body.append(
            [
                Paragraph(r.intervention, styles["Cell"]),
                Paragraph(str(r.rank), styles["Cell"]),
                Paragraph(f"{r.sucra:.3f}", styles["Cell"]),
                Paragraph(f"{r.mean_rank:.2f}", styles["Cell"]),
            ]
        )
    t = Table(body, colWidths=[2.4 * inch, 0.8 * inch, 1.0 * inch, 1.2 * inch], repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
                ("BOX", (0, 0), (-1, -1), 0.4, BORDER),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, BORDER),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


# ── PDF builder ─────────────────────────────────────────────────────────


def build_pdf(data: NmaReportData, images_dir: Path | None = None) -> bytes:
    if data.results is None:
        return b""
    buf = io.BytesIO()
    styles = make_pdf_styles()
    decor = make_page_decorations(footer_label="Network meta-analysis")
    doc = BaseDocTemplate(
        buf,
        pagesize=landscape(LETTER),
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.6 * inch,
        title=data.title,
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="page", frames=[frame], onPage=decor)])

    results = data.results
    flow: list[Any] = []
    flow.append(Paragraph("NETWORK META-ANALYSIS", styles["Eyebrow"]))
    flow.append(Paragraph(data.title, styles["Title"]))
    ts = data.generated_at.strftime("%Y-%m-%d %H:%M UTC")
    flow.append(Paragraph(f"Generated {ts}", styles["Body"]))
    flow.append(Spacer(1, 8))
    flow.append(
        Paragraph(
            f"<i>Backend: <b>{results.backend}</b>. Effect measure: "
            f"<b>{results.league_table.effect_measure}</b>. Reference: "
            f"<b>{results.pico.interventions[0]}</b>.</i>",
            styles["Body"],
        )
    )

    # PICO
    flow.append(Spacer(1, 12))
    flow.append(Paragraph("Research question + interventions", styles["H2"]))
    flow.append(
        Paragraph(
            f"<b>Population:</b> {results.pico.population}<br/>"
            f"<b>Interventions:</b> {', '.join(results.pico.interventions)}<br/>"
            f"<b>Outcome:</b> {results.pico.outcome}",
            styles["Body"],
        )
    )

    # League table
    flow.append(Spacer(1, 14))
    flow.append(Paragraph("League table — all-vs-all pairwise effects", styles["H2"]))
    flow.append(
        Paragraph(
            "<i>Read row vs column. Green cell: row protective (CI excludes "
            "null on the favourable side). Red cell: row risk-increasing. "
            "Cells reflect the chosen effect measure.</i>",
            styles["Body"],
        )
    )
    flow.append(_league_pdf(results, styles))

    # SUCRA
    flow.append(Spacer(1, 14))
    flow.append(Paragraph("SUCRA ranking", styles["H2"]))
    flow.append(
        Paragraph(
            "<i>Higher SUCRA = higher rank. SUCRA is the surface under "
            "the cumulative ranking curve; 1.0 means certainly best, "
            "0.0 means certainly worst.</i>",
            styles["Body"],
        )
    )
    flow.append(_sucra_pdf(results, styles))

    # Network plot
    resolved = _resolve_image(images_dir, results.network.image_url)
    if resolved is not None:
        flow.append(Spacer(1, 14))
        flow.append(Paragraph("Network geometry", styles["H2"]))
        flow.append(Image(str(resolved), width=8.0 * inch, height=5.5 * inch))
    elif results.network.image_url:
        flow.append(Spacer(1, 14))
        flow.append(Paragraph("Network geometry", styles["H2"]))
        flow.append(
            Paragraph(
                "<i>Network geometry image not available on this server.</i>",
                styles["Body"],
            )
        )

    # Interpretation
    flow.append(Spacer(1, 14))
    flow.append(Paragraph("Interpretation", styles["H2"]))
    flow.append(Paragraph(results.interpretation, styles["Body"]))

    if results.caveats:
        flow.append(Spacer(1, 10))
        flow.append(Paragraph("Caveats", styles["H2"]))
        for c in results.caveats:
            flow.append(Paragraph(f"• {c}", styles["Body"]))

    if results.studies_included:
        flow.append(Spacer(1, 10))
        flow.append(
            Paragraph(
                f"<i>Studies included: {', '.join(results.studies_included[:30])}"
                f"{' …' if len(results.studies_included) > 30 else ''}</i>",
                styles["Body"],
            )
        )

    doc.build(flow)
    return buf.getvalue()


# ── DOCX builder ────────────────────────────────────────────────────────


def build_docx(data: NmaReportData, images_dir: Path | None = None) -> bytes:
    if data.results is None:
        return b""
    docx = Document()
    docx.core_properties.title = data.title
    docx.add_heading(data.title, level=0)
    runs = docx.add_paragraph(
        f"Generated {data.generated_at.strftime('%Y-%m-%d %H:%M UTC')}"
    ).runs
    if runs:
        runs[0].font.color.rgb = DOCX_MUTED

    results = data.results
    p = docx.add_paragraph()
    p.add_run("Backend: ").bold = True
    p.add_run(results.backend)
    p2 = docx.add_paragraph()
    p2.add_run("Effect measure: ").bold = True
    p2.add_run(results.league_table.effect_measure)
    p3 = docx.add_paragraph()
    p3.add_run("Reference: ").bold = True
    p3.add_run(results.pico.interventions[0])

    docx.add_heading("Research question + interventions", level=1)
    for label, value in (
        ("Population", results.pico.population),
        ("Interventions", ", ".join(results.pico.interventions)),
        ("Outcome", results.pico.outcome),
    ):
        p = docx.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(value)

    docx.add_heading("League table", level=1)
    interventions = results.pico.interventions
    by_pair = {(r.row_intervention, r.col_intervention): r for r in results.league_table.rows}
    cols = [""] + list(interventions)
    t = docx.add_table(rows=1 + len(interventions), cols=len(cols))
    t.style = "Light Grid Accent 1"
    for j, c in enumerate(cols):
        cell = t.cell(0, j)
        cell.text = c
        cell.paragraphs[0].runs[0].bold = True
    for i, row_name in enumerate(interventions, start=1):
        t.cell(i, 0).text = row_name
        for j, col_name in enumerate(interventions, start=1):
            if row_name == col_name:
                t.cell(i, j).text = "—"
                continue
            cell = by_pair.get((row_name, col_name))
            if cell is None:
                t.cell(i, j).text = "—"
            else:
                t.cell(i, j).text = _fmt_effect(cell.effect, cell.ci_lower, cell.ci_upper)

    docx.add_heading("SUCRA ranking", level=1)
    cols = ["Intervention", "Rank", "SUCRA", "Mean rank"]
    sorted_rows = sorted(results.sucra, key=lambda r: r.rank)
    st = docx.add_table(rows=1 + len(sorted_rows), cols=len(cols))
    st.style = "Light Grid Accent 1"
    for j, c in enumerate(cols):
        cell = st.cell(0, j)
        cell.text = c
        cell.paragraphs[0].runs[0].bold = True
    for i, r in enumerate(sorted_rows, start=1):
        st.cell(i, 0).text = r.intervention
        st.cell(i, 1).text = str(r.rank)
        st.cell(i, 2).text = f"{r.sucra:.3f}"
        st.cell(i, 3).text = f"{r.mean_rank:.2f}"

    resolved = _resolve_image(images_dir, results.network.image_url)
    if resolved is not None:
        docx.add_heading("Network geometry", level=1)
        docx.add_picture(str(resolved), width=Inches(6.0))

    docx.add_heading("Interpretation", level=1)
    docx.add_paragraph(results.interpretation)
    if results.caveats:
        docx.add_heading("Caveats", level=1)
        for c in results.caveats:
            docx.add_paragraph(c, style="List Bullet")

    out = io.BytesIO()
    docx.save(out)
    return out.getvalue()


__all__ = [
    "NmaReportData",
    "assemble_report_data",
    "build_docx",
    "build_pdf",
]
